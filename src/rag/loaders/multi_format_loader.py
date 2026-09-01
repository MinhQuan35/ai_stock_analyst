"""
Multi-Format Document Loader
Supports: PDF, Word, Excel, CSV, Markdown, Text
"""
from pathlib import Path
from typing import List
from langchain_core.documents import Document

from src.utils import logger
from src.exceptions import DocumentLoadError


class MultiFormatLoader:
    """Load documents from multiple file formats."""
    
    SUPPORTED_EXTENSIONS = {
        ".md": "_load_text",
        ".txt": "_load_text",
        ".pdf": "_load_pdf",
        ".docx": "_load_docx",
        ".xlsx": "_load_xlsx",
        ".xls": "_load_xlsx",
        ".csv": "_load_csv",
    }
    
    def __init__(self, directory: str):
        """Initialize loader."""
        self.directory = Path(directory)
    
    def load(self) -> List[Document]:
        """Load all supported documents."""
        if not self.directory.exists():
            raise DocumentLoadError(f"Directory not found: {self.directory}")
        
        documents = []
        files = [f for f in self.directory.rglob("*") if f.is_file()]
        
        logger.info(f"Found {len(files)} files in {self.directory}")
        
        for file_path in files:
            ext = file_path.suffix.lower()
            loader_method = self.SUPPORTED_EXTENSIONS.get(ext)
            
            if not loader_method:
                logger.debug(f"Skipping unsupported file: {file_path.name}")
                continue
            
            try:
                method = getattr(self, loader_method)
                docs = method(file_path)
                documents.extend(docs)
                logger.info(f"Loaded {len(docs)} pages from {file_path.name}")
            except Exception as e:
                logger.error(f"Failed to load {file_path.name}: {e}")
        
        logger.info(f"Total documents loaded: {len(documents)}")
        return documents
    
    def _load_text(self, file_path: Path) -> List[Document]:
        """Load .md or .txt files."""
        content = file_path.read_text(encoding="utf-8")
        return [Document(
            page_content=content,
            metadata={"source": str(file_path), "type": file_path.suffix[1:]},
        )]
    
    def _load_pdf(self, file_path: Path) -> List[Document]:
        """Load .pdf files (one Document per page)."""
        from pypdf import PdfReader
        
        reader = PdfReader(str(file_path))
        documents = []
        
        for i, page in enumerate(reader.pages):
            text = page.extract_text()
            if text.strip():
                documents.append(Document(
                    page_content=text,
                    metadata={
                        "source": str(file_path),
                        "type": "pdf",
                        "page": i + 1,
                        "total_pages": len(reader.pages),
                    },
                ))
        
        return documents
    
    def _load_docx(self, file_path: Path) -> List[Document]:
        """Load .docx files."""
        from docx import Document as DocxDocument
        
        doc = DocxDocument(str(file_path))
        
        # Extract paragraphs
        paragraphs = []
        for para in doc.paragraphs:
            if para.text.strip():
                paragraphs.append(para.text)
        
        # Extract tables
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text for cell in row.cells)
                if row_text.strip():
                    paragraphs.append(row_text)
        
        content = "\n\n".join(paragraphs)
        
        return [Document(
            page_content=content,
            metadata={
                "source": str(file_path),
                "type": "docx",
                "paragraphs": len(doc.paragraphs),
            },
        )]
    
    def _load_xlsx(self, file_path: Path) -> List[Document]:
        """Load .xlsx files (one Document per sheet)."""
        from openpyxl import load_workbook
        
        wb = load_workbook(str(file_path), data_only=True, read_only=True)
        documents = []
        
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = []
            
            for row in ws.iter_rows(values_only=True):
                # Convert all cells to string
                row_str = " | ".join(str(cell) if cell is not None else "" for cell in row)
                if row_str.strip():
                    rows.append(row_str)
            
            if rows:
                content = "\n".join(rows)
                documents.append(Document(
                    page_content=content,
                    metadata={
                        "source": str(file_path),
                        "type": "xlsx",
                        "sheet": sheet_name,
                    },
                ))
        
        wb.close()
        return documents
    
    def _load_csv(self, file_path: Path) -> List[Document]:
        """Load .csv files."""
        import csv
        
        rows = []
        with open(file_path, "r", encoding="utf-8") as f:
            reader = csv.reader(f)
            for row in reader:
                if any(cell.strip() for cell in row):
                    rows.append(" | ".join(row))
        
        content = "\n".join(rows)
        
        return [Document(
            page_content=content,
            metadata={"source": str(file_path), "type": "csv"},
        )]
