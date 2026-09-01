"""
Test multi-format loader
"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent))

from src.rag.loaders.multi_format_loader import MultiFormatLoader
from src.rag.splitters.recursive_splitter import RecursiveTextSplitter
from src.utils import logger


def create_sample_files():
    """Create sample test files."""
    test_dir = Path(__file__).parent.parent / "data" / "test_files"
    test_dir.mkdir(parents=True, exist_ok=True)
    
    # Sample CSV
    csv_path = test_dir / "stocks.csv"
    csv_path.write_text("""Symbol,Price,PE,EPS
VNM,75000,17.9,4200
VIC,45000,30,1500
HPG,25000,8,3125
VCB,95000,15,6333
""", encoding="utf-8")
    
    # Sample DOCX
    from docx import Document as DocxDocument
    docx_path = test_dir / "report.docx"
    doc = DocxDocument()
    doc.add_heading("Stock Analysis Report", 0)
    doc.add_paragraph("P/E ratio measures valuation.")
    doc.add_paragraph("RSI above 70 means overbought.")
    doc.add_table(rows=2, cols=2)
    doc.save(str(docx_path))
    
    # Sample XLSX
    from openpyxl import Workbook
    xlsx_path = test_dir / "portfolio.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Holdings"
    ws.append(["Symbol", "Shares", "Avg Price"])
    ws.append(["VNM", 100, 70000])
    ws.append(["FPT", 50, 130000])
    wb.save(str(xlsx_path))
    
    # Sample PDF
    from pypdf import PdfWriter
    pdf_path = test_dir / "policy.pdf"
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    # Note: Creating blank PDF without content
    with open(pdf_path, "wb") as f:
        writer.write(f)
    
    logger.info(f"Created sample files in {test_dir}")
    return test_dir


def main():
    print("=" * 50)
    print("  Multi-Format Loader Test")
    print("=" * 50)
    
    # Create samples
    test_dir = create_sample_files()
    
    # Load all
    print(f"\n[1] Loading from: {test_dir}")
    loader = MultiFormatLoader(str(test_dir))
    documents = loader.load()
    
    print(f"\n[2] Loaded {len(documents)} documents:")
    for i, doc in enumerate(documents):
        print(f"  [{i+1}] {doc.metadata.get('type', 'unknown')}: {doc.page_content[:80]}...")
    
    # Split
    print(f"\n[3] Splitting documents...")
    splitter = RecursiveTextSplitter(chunk_size=200, chunk_overlap=20)
    chunks = splitter.split(documents)
    print(f"  Created {len(chunks)} chunks")
    
    print("\n" + "=" * 50)
    print("  TEST COMPLETE!")
    print("=" * 50)


if __name__ == "__main__":
    main()
